from __future__ import annotations

import io
import zipfile

import fitz
import pytest
from docx import Document
from openpyxl import Workbook

from mote.contracts.file import DocumentResourceLimitError, SearchSkipReason
from mote.runtime.fileops import documents as documents_module
from mote.runtime.fileops import text_sources as text_sources_module
from mote.runtime.fileops.documents import ExtractionBudget, extract_document_bytes, extract_pdf_text
from mote.ztest.fileops_factory import FileOperations

_LARGE_ARCHIVE_ENTRY_BYTES = 8_192
_ARCHIVE_LIMIT_BYTES = 1_024
_OUTPUT_LIMIT_BYTES = 4
_GENEROUS_ARCHIVE_LIMIT_BYTES = 16 * 1_024 * 1_024


def _budget(*, archive: int, output: int) -> ExtractionBudget:
    return ExtractionBudget(
        max_archive_uncompressed_bytes=archive,
        max_output_bytes=output,
    )


def _compressed_archive(entry_name: str, payload: bytes) -> bytes:
    stream = io.BytesIO()
    with zipfile.ZipFile(
        stream,
        mode="w",
        compression=zipfile.ZIP_DEFLATED,
    ) as archive:
        archive.writestr(entry_name, payload)
    return stream.getvalue()


def _docx_bytes(text: str) -> bytes:
    document = Document()
    document.add_paragraph(text)
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def _xlsx_bytes(text: str) -> bytes:
    workbook = Workbook()
    workbook.active["A1"] = text
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def _pdf_bytes(text: str) -> bytes:
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), text)
    raw = document.tobytes()
    document.close()
    return raw


@pytest.mark.parametrize(
    ("suffix", "entry_name"),
    [
        (".docx", "word/document.xml"),
        (".xlsx", "xl/worksheets/sheet1.xml"),
    ],
)
def test_zip_documents_reject_declared_uncompressed_size_before_extraction(
    suffix,
    entry_name,
):
    raw = _compressed_archive(entry_name, b"x" * _LARGE_ARCHIVE_ENTRY_BYTES)
    assert len(raw) < _ARCHIVE_LIMIT_BYTES

    with pytest.raises(DocumentResourceLimitError) as exc_info:
        extract_document_bytes(
            raw,
            suffix,
            budget=_budget(
                archive=_ARCHIVE_LIMIT_BYTES,
                output=_GENEROUS_ARCHIVE_LIMIT_BYTES,
            ),
        )

    assert exc_info.value.context == {
        "resource": "archive_uncompressed_bytes",
        "consumed": _LARGE_ARCHIVE_ENTRY_BYTES,
        "maximum": _ARCHIVE_LIMIT_BYTES,
    }


@pytest.mark.parametrize(
    ("suffix", "raw"),
    [
        (".pdf", _pdf_bytes("oversized PDF output")),
        (".docx", _docx_bytes("oversized DOCX output")),
        (".xlsx", _xlsx_bytes("oversized XLSX output")),
    ],
)
def test_document_output_budget_has_one_typed_failure_contract(suffix, raw):
    with pytest.raises(DocumentResourceLimitError) as exc_info:
        extract_document_bytes(
            raw,
            suffix,
            budget=_budget(
                archive=_GENEROUS_ARCHIVE_LIMIT_BYTES,
                output=_OUTPUT_LIMIT_BYTES,
            ),
        )

    assert exc_info.value.context["resource"] == "output_bytes"
    assert exc_info.value.context["consumed"] > _OUTPUT_LIMIT_BYTES
    assert exc_info.value.context["maximum"] == _OUTPUT_LIMIT_BYTES
    assert set(exc_info.value.context) == {"resource", "consumed", "maximum"}


def test_pdf_budget_failure_never_falls_through_to_another_adapter(
    tmp_path,
    monkeypatch,
):
    target = tmp_path / "bounded.pdf"
    target.write_bytes(b"not parsed by the stub adapters")
    fallback_called = False

    def exhausted(_path, *, sink):
        sink.write("x" * (_OUTPUT_LIMIT_BYTES + 1))

    def fallback(_path, *, sink):
        nonlocal fallback_called
        fallback_called = True
        sink.write("must not run")

    monkeypatch.setattr(
        documents_module,
        "_PDF_ADAPTERS",
        (exhausted, fallback),
    )

    with pytest.raises(DocumentResourceLimitError):
        extract_pdf_text(
            str(target),
            budget=_budget(
                archive=_GENEROUS_ARCHIVE_LIMIT_BYTES,
                output=_OUTPUT_LIMIT_BYTES,
            ),
        )

    assert not fallback_called


def test_search_projects_document_resource_limits_without_losing_the_reason(
    tmp_path,
    monkeypatch,
):
    project = tmp_path / "project"
    project.mkdir()
    target = project / "bounded.docx"
    target.write_bytes(b"sealed document bytes")
    operations = FileOperations(
        session_id="document-budget-search",
        journal_path=tmp_path / "state" / "rollout.jsonl",
        get_project_root=lambda: str(project),
        lock_root=tmp_path / "locks",
    )

    def exhausted(_raw, _suffix, *, budget):
        raise DocumentResourceLimitError(
            "document extraction exceeded its output budget",
            resource="output_bytes",
            consumed=budget.max_output_bytes + 1,
            maximum=budget.max_output_bytes,
        )

    monkeypatch.setattr(text_sources_module, "extract_document_bytes", exhausted)

    result = operations.search(root=str(target), content="needle")

    assert result.rows == ()
    assert result.summary.skipped_files == 1
    assert result.skipped[0].reason == SearchSkipReason.RESOURCE_LIMIT
    assert "output budget" in result.skipped[0].detail
